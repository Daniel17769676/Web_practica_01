from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login
from django.contrib import messages
from reservas.models import Reserva
from servicios.models import Disponibilidad
from usuarios.models import Administrador
from django.contrib.auth.decorators import login_required
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from django.utils import timezone
from datetime import datetime
from openpyxl.styles import PatternFill


# Decorador personalizado para verificar sesión de admin
def admin_login_required(view_func):
    def _wrapped_view(request, *args, **kwargs):
        if not request.session.get('is_admin_logged_in'):
            messages.error(request, 'Debe iniciar sesión como administrador')
            return redirect('base')
        return view_func(request, *args, **kwargs)
    return _wrapped_view

@admin_login_required  # Usa SOLO tu decorador personalizado
def panel_administrador(request):
    # Obtener datos del admin desde la sesión
    admin_id = request.session.get('admin_id')
    
    try:
        admin = Administrador.objects.get(id=admin_id)
    except Administrador.DoesNotExist:
        messages.error(request, 'Administrador no encontrado')
        return redirect('base')
    
    # ✅ CONSULTA DE RESERVAS
    consulta_admin = Reserva.objects.select_related(
        'disponibilidad__servicio'
    ).order_by('-fecha_reserva', '-horario')

    if not consulta_admin.exists():
        messages.info(request, 'No hay reservas registradas.')
    
    # Debug para verificar
    print(f"Reservas encontradas: {consulta_admin.count()}")
    for reserva in consulta_admin:
        print(f"Reserva ID: {reserva.id}, Cliente: {reserva.cliente_nombre}")
    
    context = {
        'admin': admin,
        'reservas': consulta_admin
    }
    
    return render(request, 'usuarios/panel_administrador.html', context)

def login_administrador(request):
    if request.method == 'POST':
        email = request.POST.get('email')
        password = request.POST.get('password')
        
        try:
            # Buscar el administrador por email
            admin = Administrador.objects.get(correo=email)
            
            # 🔐 Verificar contraseña usando el nuevo método
            if admin.check_password(password):
                # Crear variables de sesión
                request.session['admin_id'] = admin.id
                request.session['admin_nombre'] = admin.nombre
                request.session['admin_email'] = admin.correo
                request.session['is_admin_logged_in'] = True
                
                messages.success(request, f'¡Bienvenido {admin.nombre}!')
                return redirect('usuarios:panel_administrador')
            else:
                messages.error(request, 'Contraseña incorrecta')
                
        except Administrador.DoesNotExist:
            messages.error(request, 'No existe un administrador con este email')

    return render(request, 'base.html')

def logout_administrador(request):
    # Limpiar la sesión
    request.session.flush()
    messages.success(request, 'Sesión cerrada correctamente')
    return redirect('base')

#View que ELIMINA LA RESERVA (Solicitud de cancelación)
def cancelar_reserva(request, reserva_id):
    if request.method == 'DELETE':
        try:
            reserva = Reserva.objects.get(id=reserva_id)
            reserva.delete()
            # Respuesta JSON para AJAX
            return JsonResponse({
                'status': 'success',
                'message': 'Reserva cancelada correctamente'
            })
        except Reserva.DoesNotExist:
            return JsonResponse({
                'status': 'error',
                'message': 'Reserva no encontrada'
            }, status=404)

    return JsonResponse({'status': 'error', 'message': 'Método no permitido'}, status=405)

#View que CONFIRMA LA RESERVA (Rechaza la solicitud de CANCELACION)
def confirmar_reserva(request, reserva_id):
    if request.method == 'POST':
        try:
            reserva = Reserva.objects.get(id=reserva_id)
            reserva.estado = 'confirmada'
            reserva.save()

            # Respuesta JSON para AJAX
            return JsonResponse({
                'status': 'success',
                'message': 'Reserva confirmada correctamente'
            })
            
        except Reserva.DoesNotExist:
            return JsonResponse({
                'status': 'error',
                'message': 'Reserva no encontrada'
            }, status=404)
    
    return JsonResponse({'status': 'error', 'message': 'Método no permitido'}, status=405)

def generar_reportes(request):
    if request.method == 'POST':
        # Obtener parámetros del formulario
        tipo_reporte = request.POST.get('reporte-tipo')
        fecha_desde = request.POST.get('reporte-desde')
        fecha_hasta = request.POST.get('reporte-hasta')
        
        # Crear libro de Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte Reservas"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center")
        
        if tipo_reporte == 'reservas':
            # Reporte de reservas por fecha
            queryset = Reserva.objects.filter(
                fecha_reserva__range=[fecha_desde, fecha_hasta]
            )
            
            # Encabezados
            headers = ['ID', 'Cliente', 'Servicio', 'Fecha', 'Hora', 'Estado']
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = alignment
            
            # Datos
            for row_num, reserva in enumerate(queryset, 2):
                ws.cell(row=row_num, column=1).value = reserva.id
                ws.cell(row=row_num, column=2).value = reserva.cliente_nombre
                ws.cell(row=row_num, column=3).value = reserva.disponibilidad.servicio.nombre if reserva.disponibilidad and reserva.disponibilidad.servicio else 'N/A'
                ws.cell(row=row_num, column=4).value = reserva.fecha_reserva.strftime('%d/%m/%Y')
                ws.cell(row=row_num, column=5).value = reserva.horario
                ws.cell(row=row_num, column=6).value = reserva.estado
                
        
        elif tipo_reporte == 'servicios':

            # Reporte de servicios
            queryset = Disponibilidad.objects.select_related('servicio').all()

            # Encabezados
            headers = ['ID Servicio', 'Administrador', 'Hora_inicio', 'Hora_Fin', 'Duracion_servicio', 'Ubicacion', 'Fecha_fin', 'Fecha_inicio']
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = alignment

            # Datos
            for row_num, disponibilidad in enumerate(queryset, 2):
                ws.cell(row=row_num, column=1).value = disponibilidad.servicio.id
                ws.cell(row=row_num, column=2).value = disponibilidad.administrador if disponibilidad.administrador else 'N/A'
                ws.cell(row=row_num, column=3).value = disponibilidad.hora_inicio.strftime('%H:%M')
                ws.cell(row=row_num, column=4).value = disponibilidad.hora_fin.strftime('%H:%M')
                ws.cell(row=row_num, column=5).value = disponibilidad.intervalo
                ws.cell(row=row_num, column=6).value = disponibilidad.ubicacion
                ws.cell(row=row_num, column=7).value = disponibilidad.fecha_fin.strftime('%d/%m/%Y') if disponibilidad.fecha_fin else 'N/A'
                ws.cell(row=row_num, column=8).value = disponibilidad.fecha_inicio.strftime('%d/%m/%Y') if disponibilidad.fecha_inicio else 'N/A'

        # Ajustar anchos de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Crear respuesta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename=reporte_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        
        wb.save(response)
        return response
    
    return HttpResponse("Método no permitido", status=405)
